import ctypes
import os
import sys
import re
import pandas as pd
from xml.sax.saxutils import escape

try:
    from docx import Document
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    MISSING_LIBS = False
except ImportError:
    MISSING_LIBS = True

from PySide6.QtCore import Qt
from PySide6.QtGui import QFont, QIcon
from PySide6.QtWidgets import (
    QApplication,
    QButtonGroup,
    QFileDialog,
    QGroupBox,
    QHBoxLayout,
    QLabel,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QRadioButton,
    QVBoxLayout,
    QWidget,
)

try:
    ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID(
        "granola.exporter.gui.1.0"
    )
except Exception:
    pass

def resource_path(relative_path):
    if hasattr(sys, "_MEIPASS"):
        return os.path.join(sys._MEIPASS, relative_path)
    return os.path.join(os.path.abspath("."), relative_path)

class GranolaExporterApp(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Granola Notes Exporter")
        self.resize(500, 320)
        
        self.setStyleSheet("""
            * {
                font-family: "Martel Sans", ".AppleSystemUIFont", "Helvetica Neue", Arial, sans-serif;
            }
            QMainWindow, QMessageBox { background-color: #e9eaf0; }
            QLabel { font-size: 16px; color: #113859; font-weight: bold; }
            QGroupBox { 
                font-weight: bold; 
                border: 2px solid #97CAEB; 
                border-radius: 6px; 
                margin-top: 15px; 
            }
            QGroupBox::title { 
                subcontrol-origin: margin; 
                left: 10px; 
                padding: 0 5px 0 5px; 
                color: #113859;
            }
            
            QRadioButton { 
                font-size: 13px; 
                color: #113859; 
                padding: 5px; 
            }
            
            QRadioButton::indicator {
                width: 14px;
                height: 14px;
                border-radius: 8px; 
                border: 2px solid #696969;
                background-color: #e9eaf0; 
            }

            QRadioButton::indicator:checked {
                border: 2px solid #113859;
                background-color: #F78D2D;
            }

            QPushButton { 
                background-color: #113859; 
                color: #e9eaf0; 
                border: 1px solid #113859; 
                border-radius: 4px; 
                padding: 6px 12px; 
                font-weight: bold;
            }
            QPushButton:hover { 
                background-color: #696969;
                color: #e9eaf0; 
                border: 1px solid #696969;
            }

            #exportButton { 
                background-color: #113859;
                color: #e9eaf0; 
                font-weight: bold; 
                font-size: 15px; 
                border: none;
                border-radius: 6px; 
            }
            #exportButton:hover { 
                background-color: #696969; 
                color: #e9eaf0; 
            }
        """)
        
        self.csv_path = None
        self.central_widget = QWidget()
        self.setCentralWidget(self.central_widget)
        self.layout = QVBoxLayout(self.central_widget)
        self.layout.setSpacing(15)
        self.layout.setContentsMargins(25, 25, 25, 25)

        self.init_ui()

        if MISSING_LIBS:
            QMessageBox.critical(self, "Missing Libraries", 
                "Missing python-docx or reportlab.\nPlease install them to use Word and PDF exports.")

    def init_ui(self):
        lbl_header = QLabel(" Granola Notes Export Parser")
        font = QFont()
        font.setPointSize(18)
        font.setBold(True)
        lbl_header.setFont(font)
        self.layout.addWidget(lbl_header)

        file_layout = QHBoxLayout()
        self.lbl_file = QLabel("No CSV file selected...")
        self.lbl_file.setStyleSheet("color: #6c757d; font-style: italic;")
        
        btn_browse = QPushButton("Browse")
        btn_browse.clicked.connect(self.browse_file)
        
        file_layout.addWidget(self.lbl_file)
        file_layout.addStretch()
        file_layout.addWidget(btn_browse)
        self.layout.addLayout(file_layout)

        group_box = QGroupBox("Select Output Format")
        group_layout = QVBoxLayout()
        group_layout.setSpacing(10)
        
        self.radio_group = QButtonGroup(self)
        self.rb_markdown = QRadioButton("Markdown (.md) - Best for note apps")
        self.rb_markdown.setChecked(True)
        self.rb_word = QRadioButton("Word Document (.docx) - Best for editing")
        self.rb_pdf = QRadioButton("PDF (.pdf) - Best for sharing")
        
        if MISSING_LIBS:
            self.rb_word.setEnabled(False)
            self.rb_pdf.setEnabled(False)
            
        self.radio_group.addButton(self.rb_markdown, 1)
        self.radio_group.addButton(self.rb_word, 2)
        self.radio_group.addButton(self.rb_pdf, 3)
        
        group_layout.addWidget(self.rb_markdown)
        group_layout.addWidget(self.rb_word)
        group_layout.addWidget(self.rb_pdf)
        group_box.setLayout(group_layout)
        self.layout.addWidget(group_box)

        self.layout.addStretch()

        self.btn_export = QPushButton("Export Notes")
        self.btn_export.setObjectName("exportButton")
        self.btn_export.setMinimumHeight(45)
        self.btn_export.clicked.connect(self.run_export)
        self.layout.addWidget(self.btn_export)

        lbl_about = QLabel(
            "Version 1.0.0 | Vibe coded with <span style='color: #F78D2D; font-size: 13px;'>♥</span> by "
            "<a href='https://www.linkedin.com/in/sean-m-walsh/' style='color: #696969; text-decoration: none;'>Sean Walsh</a>"
        )
        lbl_about.setOpenExternalLinks(True)
        lbl_about.setAlignment(Qt.AlignCenter)
        lbl_about.setStyleSheet("""
            font-size: 11px; 
            color: #696969; 
            font-weight: normal;
            margin-top: 10px;
        """)
        self.layout.addWidget(lbl_about)

    def browse_file(self):
        file_name, _ = QFileDialog.getOpenFileName(self, "Select Granola Export CSV", "", "CSV Files (*.csv)")
        if file_name:
            self.csv_path = file_name
            self.lbl_file.setText(f"📁 {os.path.basename(file_name)}")
            self.lbl_file.setStyleSheet("color: #6c757d; font-weight: bold;")

    def run_export(self):
        if not self.csv_path:
            QMessageBox.warning(self, "No File", "Please select a CSV file first.")
            return

        choice_id = self.radio_group.checkedId()
        format_map = {1: ('.md', 'granola_markdown'), 2: ('.docx', 'granola_word_docs'), 3: ('.pdf', 'granola_pdfs')}
        ext, folder_name = format_map[choice_id]
        
        base_dir = os.path.dirname(self.csv_path)
        output_dir = os.path.join(base_dir, folder_name)
        os.makedirs(output_dir, exist_ok=True)

        try:
            df = pd.read_csv(self.csv_path)
            for _, row in df.iterrows():
                title = str(row['document_title'])
                date = str(row['document_created'])[:10]
                workspace = str(row['workspace_name'])
                summary = str(row['summary'])
                transcript = str(row['transcript'])
                
                safe_title = re.sub(r'[\\/*?:"<>|]', "", title).strip()
                filename = os.path.join(output_dir, f"{date}_{safe_title}{ext}")
                
                if choice_id == 1:
                    content = f"# {title}\n**Date:** {date}\n**Workspace:** {workspace}\n\n## Summary\n{summary}\n\n## Transcript\n{transcript}"
                    with open(filename, 'w', encoding='utf-8') as f:
                        f.write(content)
                        
                elif choice_id == 2:
                    doc = Document()
                    doc.add_heading(title, level=1)
                    doc.add_paragraph(f"Date: {date}\nWorkspace: {workspace}")
                    doc.add_heading("Summary", level=2)
                    doc.add_paragraph(summary)
                    doc.add_heading("Transcript", level=2)
                    doc.add_paragraph(transcript)
                    doc.save(filename)
                    
                elif choice_id == 3:
                    doc = SimpleDocTemplate(filename, pagesize=letter)
                    styles = getSampleStyleSheet()
                    story = []
                    story.append(Paragraph(f"<b>Title:</b> {escape(title)}", styles['Heading1']))
                    story.append(Paragraph(f"<b>Date:</b> {escape(date)}<br/><b>Workspace:</b> {escape(workspace)}", styles['Normal']))
                    story.append(Spacer(1, 12))
                    story.append(Paragraph("<b>Summary</b>", styles['Heading2']))
                    for line in summary.split('\n'):
                        if line.strip():
                            story.append(Paragraph(escape(line.strip()), styles['Normal']))
                    story.append(Spacer(1, 12))
                    story.append(Paragraph("<b>Transcript</b>", styles['Heading2']))
                    for line in transcript.split('\n'):
                        if line.strip():
                            story.append(Paragraph(escape(line.strip()), styles['Normal']))
                    doc.build(story)

            success_msg = (
                f"Successfully exported <b>{len(df)}</b> files to:<br><br>"
                f"<span style='color: #696969; font-size: 14px;'>"
                f"{output_dir}</span>"
            )
            QMessageBox.information(self, "Success", success_msg)
        except Exception as e:
            error_msg = (
                f"Error exporting <b>{len(df)}</b> files to:<br><br>"
                f"<span style='color: #696969; font-size: 14px;'>"
                f"{output_dir}</span>"
            )
            QMessageBox.critical(self, "Error", error_msg)


if __name__ == "__main__":
    app = QApplication(sys.argv)
    
    app.setWindowIcon(QIcon(resource_path("icon.png")))
    
    window = GranolaExporterApp()
    window.show()
    
    sys.exit(app.exec())